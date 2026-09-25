"""
Turning a paper into a Word document you can actually read.

  "the documents he makes must be strictly in word only as that's what I view
   things in"

So .docx is the deliverable, not a side-export of a Markdown file. That has
consequences beyond the file extension: fifty pages is only navigable if it has
real Heading styles (which is what drives Word's navigation pane and any table
of contents), so the structure has to be built with styles rather than with
bold text that merely looks like a heading.

Everything degrades rather than fails. If python-docx is missing the text is
still written, as Markdown, and the caller is told plainly — losing a research
run because a formatting library is absent would be an absurd way to fail.
"""

from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any

from .constants import BASE_DIR

#: Where research lands, unless a caller says otherwise.
DEFAULT_DIR = BASE_DIR / "config" / "research"


def available() -> bool:
    try:
        import docx  # noqa: F401
        return True
    except Exception:
        return False


def _safe_name(topic: str) -> str:
    cleaned = re.sub(r"[^\w\s-]", "", str(topic or "research")).strip()
    cleaned = re.sub(r"[\s_]+", "-", cleaned)[:60].strip("-")
    return cleaned.lower() or "research"


def _add_body(document: Any, text: str) -> None:
    """Write a section body, honouring subheadings and paragraph breaks.

    The model is asked for subheadings where material divides, so they have to
    survive into the document as real Heading 3s — otherwise a 1,300-word
    section arrives as one unbroken wall of text, which is exactly what makes
    long documents unreadable.
    """
    for block in str(text or "").split("\n\n"):
        block = block.strip()
        if not block:
            continue
        heading = re.match(r"^#{1,6}\s+(.*)$", block)
        if heading:
            document.add_heading(heading.group(1).strip()[:200], level=3)
            continue
        # A short bold-only line is a subheading in all but name.
        bold = re.match(r"^\*\*(.+?)\*\*:?$", block)
        if bold and len(bold.group(1)) < 90:
            document.add_heading(bold.group(1).strip(), level=3)
            continue
        if block.startswith(("- ", "* ", "• ")):
            for line in block.splitlines():
                line = re.sub(r"^[\-\*•]\s*", "", line.strip())
                if line:
                    document.add_paragraph(_plain(line), style="List Bullet")
            continue
        if re.match(r"^\d+[\.\)]\s", block):
            for line in block.splitlines():
                line = re.sub(r"^\d+[\.\)]\s*", "", line.strip())
                if line:
                    document.add_paragraph(_plain(line), style="List Number")
            continue
        document.add_paragraph(_plain(block))


def _plain(text: str) -> str:
    """Strip Markdown emphasis. Word has its own; leaving asterisks in a
    finished document looks like a bug, because it is one."""
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\s)(.+?)(?<!\s)\*(?!\*)", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text.strip()


def write_paper(paper: Any, folder: Path | None = None) -> tuple[Path | None, str]:
    """Write *paper* as .docx. Returns (path, note); path is None on failure."""
    folder = Path(folder) if folder else DEFAULT_DIR
    try:
        folder.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        return None, f"could not create {folder}: {exc}"

    stem = f"{datetime.now():%Y-%m-%d}-{_safe_name(paper.topic)}"
    if not available():
        fallback = folder / f"{stem}.md"
        try:
            fallback.write_text(_as_markdown(paper), encoding="utf-8")
        except OSError as exc:
            return None, f"could not write the paper: {exc}"
        return fallback, ("python-docx is not installed, so this was saved as "
                          "Markdown instead of Word — say 'install python-docx' "
                          "and I'll fix that.")

    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    document = docx.Document()
    for name, size in (("Normal", 11), ("Title", 28)):
        try:
            document.styles[name].font.size = Pt(size)
        except Exception:
            pass

    from .deep_research import STYLES

    document.add_heading(str(paper.topic)[:200], level=0)
    subtitle = document.add_paragraph(
        f"{STYLES[paper.style].label} · {paper.pages} pages · "
        f"{paper.words:,} words")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    stamp = document.add_paragraph(
        f"Compiled by O.R.I.O.N. — {datetime.now():%d %B %Y, %H:%M}")
    stamp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # The scholarly frame: abstract and contributions up front, as a paper has.
    if getattr(paper, "abstract", ""):
        document.add_page_break()
        document.add_heading("Abstract", level=1)
        _add_body(document, paper.abstract)
    if getattr(paper, "contributions", None):
        document.add_heading("Key contributions", level=2)
        for point in paper.contributions:
            document.add_paragraph(_plain(point), style="List Bullet")

    # A contents list rather than a Word TOC field: a real field shows
    # "Right-click to update" until the reader does so, which is worse than a
    # plain list that is simply correct when it arrives.
    document.add_page_break()
    document.add_heading("Contents", level=1)
    for section in paper.sections:
        document.add_paragraph(f"{section.number}. {section.title}",
                               style="List Number" if False else None)

    for section in paper.sections:
        document.add_page_break()
        document.add_heading(f"{section.number}. {section.title}", level=1)
        if section.error:
            note = document.add_paragraph(
                f"[This section could not be written: {section.error}. "
                "Ask me to retry it.]")
            note.runs[0].italic = True
            continue
        _add_body(document, section.body)

    if paper.sources:
        document.add_page_break()
        document.add_heading("References", level=1)
        for index, source in enumerate(paper.sources, 1):
            title = str(source.get("title") or "untitled")
            url = str(source.get("url") or "")
            document.add_paragraph(f"{index}. {title}" + (f" — {url}" if url else ""))

    path = folder / f"{stem}.docx"
    try:
        document.save(str(path))
    except Exception as exc:
        return None, f"could not save the Word document: {exc}"
    return path, ""


def _as_markdown(paper: Any) -> str:
    lines = [f"# {paper.topic}", "",
             f"_Compiled by O.R.I.O.N. — {datetime.now():%d %B %Y}_", ""]
    for section in paper.sections:
        lines += [f"## {section.number}. {section.title}", "",
                  section.body or f"_(not written: {section.error})_", ""]
    if paper.sources:
        lines += ["## References", ""]
        for index, source in enumerate(paper.sources, 1):
            lines.append(f"{index}. {source.get('title','untitled')} "
                         f"{source.get('url','')}".strip())
    return "\n".join(lines)


__all__ = ["DEFAULT_DIR", "available", "write_paper"]
